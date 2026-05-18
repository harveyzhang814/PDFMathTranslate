"""
E2E test: full translate + Word export pipeline on a real PDF fixture.

Run with:
    python3 -m pytest test/e2e/ -v -m e2e

First run creates test/e2e/expected/word_3pages_baseline.json;
subsequent runs compare counts within ±20%.
"""
import json
from pathlib import Path

import pytest

FIXTURE_PDF = Path(__file__).parent / "fixtures" / "S0894731722001407.pdf"
EXPECTED_DIR = Path(__file__).parent / "expected"
OUTPUT_DIR = Path(__file__).parent / "output" / "word_3pages"
BASELINE_FILE = EXPECTED_DIR / "word_3pages_baseline.json"

# Pages 3-5 (0-indexed) are PDF pages 4-6, which contain figures
PAGES = [3, 4, 5]
TOLERANCE = 0.20


@pytest.fixture(scope="module")
def onnx_model():
    from pdf2zh.doclayout import OnnxModel
    return OnnxModel.load_available()


@pytest.fixture(scope="module")
def docx_path(onnx_model):
    import shutil
    import tempfile
    from pdf2zh.doclayout import ModelInstance
    from pdf2zh.export_word import export_pdf_to_word
    from pdf2zh.kernel import KernelRegistry
    from pdf2zh.kernel.protocol import TranslateRequest

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ModelInstance.value = onnx_model

    KernelRegistry.switch("fast")
    kernel = KernelRegistry.get()

    elem_dir = tempfile.mkdtemp(prefix="pdf2zh_e2e_elem_")
    try:
        request = TranslateRequest(
            files=[str(FIXTURE_PDF)],
            output=str(OUTPUT_DIR),
            lang_in="en",
            lang_out="zh",
            service="google",
            pages=PAGES,
            extract_elements=True,
            elements_output_dir=elem_dir,
        )
        results = kernel.translate(request)
        mono = Path(results[0].mono_pdf)
        docx = str(OUTPUT_DIR / f"{mono.stem}.docx")
        export_pdf_to_word(str(mono), elem_dir, docx, lang_out="zh", pages=PAGES)
    finally:
        shutil.rmtree(elem_dir, ignore_errors=True)

    return Path(docx)


@pytest.mark.e2e
class TestWordExportE2E:
    def test_docx_file_exists(self, docx_path):
        assert docx_path.exists(), f"Expected docx at {docx_path}"
        assert docx_path.suffix == ".docx"

    def test_docx_has_reasonable_size(self, docx_path):
        size_kb = docx_path.stat().st_size / 1024
        assert size_kb > 50, f"docx suspiciously small: {size_kb:.1f} KB"

    def test_docx_has_paragraphs(self, docx_path):
        import docx as python_docx
        doc = python_docx.Document(str(docx_path))
        non_empty = [p for p in doc.paragraphs if p.text.strip()]
        assert len(non_empty) >= 10, f"Too few paragraphs: {len(non_empty)}"

    def test_docx_has_images(self, docx_path):
        import docx as python_docx
        doc = python_docx.Document(str(docx_path))
        assert len(doc.inline_shapes) >= 1, "No images found in docx"

    def test_structural_regression(self, docx_path):
        """Paragraph and image counts must stay within ±20% of the saved baseline."""
        import docx as python_docx
        doc = python_docx.Document(str(docx_path))
        current = {
            "para_count": len([p for p in doc.paragraphs if p.text.strip()]),
            "img_count": len(doc.inline_shapes),
        }

        if not BASELINE_FILE.exists():
            BASELINE_FILE.parent.mkdir(parents=True, exist_ok=True)
            BASELINE_FILE.write_text(json.dumps(current, indent=2))
            # Baseline created on first run; test trivially passes.
            return

        baseline = json.loads(BASELINE_FILE.read_text())
        for key, val in current.items():
            expected = baseline[key]
            lo = expected * (1 - TOLERANCE)
            hi = expected * (1 + TOLERANCE)
            assert lo <= val <= hi, (
                f"{key}: got {val}, baseline {expected} (±{TOLERANCE:.0%} → [{lo:.0f}, {hi:.0f}])"
            )
