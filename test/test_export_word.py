import os
import tempfile
import unittest
from unittest.mock import MagicMock, patch


class TestExportPdfToWord(unittest.TestCase):
    """Tests for export_pdf_to_word in pdf2zh/export_word.py."""

    # Chinese sample texts used throughout so the CJK filter does not strip them
    _ZH = "这是翻译后的中文段落。"
    _ZH2 = "右栏中文内容。"
    _ZH_LEFT = "左栏中文内容。"

    def _make_mock_page(self, blocks, width=612.0, height=792.0):
        """Return a mock pymupdf Page with the given text blocks."""
        page = MagicMock()
        page.rect.width = width
        page.rect.height = height
        page.get_text.return_value = [
            (b["x0"], b["y0"], b["x1"], b["y1"], b["content"], 0, 0)
            for b in blocks
        ]
        return page

    @patch("pdf2zh.export_word.sort_text_blocks_by_layout")
    @patch("pymupdf.open")
    def test_sort_result_is_used(self, mock_pymupdf_open, mock_sort):
        """Bug fix: sort_text_blocks_by_layout return value must be assigned back."""
        blocks_in = [
            {"x0": 300.0, "y0": 100.0, "x1": 580.0, "y1": 120.0, "content": self._ZH2},
            {"x0": 50.0,  "y0": 100.0, "x1": 280.0, "y1": 120.0, "content": self._ZH_LEFT},
        ]
        sorted_blocks = [
            {"x0": 50.0,  "y0": 100.0, "x1": 280.0, "y1": 120.0, "content": self._ZH_LEFT},
            {"x0": 300.0, "y0": 100.0, "x1": 580.0, "y1": 120.0, "content": self._ZH2},
        ]
        mock_sort.return_value = sorted_blocks

        mock_doc = MagicMock()
        mock_page = self._make_mock_page(blocks_in)
        mock_doc.__iter__ = MagicMock(return_value=iter([mock_page]))
        mock_pymupdf_open.return_value = mock_doc

        from docx import Document as RealDocument
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            out_path = f.name
        try:
            from pdf2zh.export_word import export_pdf_to_word
            export_pdf_to_word("fake.pdf", None, out_path)

            mock_sort.assert_called_once()
            doc = RealDocument(out_path)
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            self.assertIn(self._ZH_LEFT, paras)
            self.assertIn(self._ZH2, paras)
            self.assertLess(paras.index(self._ZH_LEFT), paras.index(self._ZH2))
        finally:
            os.unlink(out_path)

    @patch("pymupdf.open")
    def test_empty_page_does_not_abort_subsequent_pages(self, mock_pymupdf_open):
        """Bug fix: empty-text page must use continue, not return, so later pages are kept."""
        page_empty = self._make_mock_page([])
        page_text  = self._make_mock_page([
            {"x0": 50.0, "y0": 200.0, "x1": 500.0, "y1": 220.0, "content": self._ZH},
        ])

        mock_doc = MagicMock()
        mock_doc.__iter__ = MagicMock(return_value=iter([page_empty, page_text]))
        mock_pymupdf_open.return_value = mock_doc

        from docx import Document as RealDocument
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            out_path = f.name
        try:
            from pdf2zh.export_word import export_pdf_to_word
            export_pdf_to_word("fake.pdf", None, out_path)

            doc = RealDocument(out_path)
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            self.assertIn(self._ZH, paras)
        finally:
            os.unlink(out_path)

    @patch("pymupdf.open")
    def test_three_pages_all_processed(self, mock_pymupdf_open):
        """All pages in a multi-page document are written to the output."""
        pages = [
            self._make_mock_page([{"x0": 50.0, "y0": 100.0, "x1": 500.0, "y1": 120.0, "content": f"第{i}页翻译内容"}])
            for i in range(1, 4)
        ]
        mock_doc = MagicMock()
        mock_doc.__iter__ = MagicMock(return_value=iter(pages))
        mock_pymupdf_open.return_value = mock_doc

        from docx import Document as RealDocument
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            out_path = f.name
        try:
            from pdf2zh.export_word import export_pdf_to_word
            export_pdf_to_word("fake.pdf", None, out_path)

            doc = RealDocument(out_path)
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            for i in range(1, 4):
                self.assertIn(f"第{i}页翻译内容", paras)
        finally:
            os.unlink(out_path)

    @patch("pymupdf.open")
    def test_pages_filter_skips_other_pages(self, mock_pymupdf_open):
        """pages parameter: only specified pages appear in the output."""
        pages = [
            self._make_mock_page([{"x0": 50.0, "y0": 100.0, "x1": 500.0, "y1": 120.0, "content": f"第{i}页内容"}])
            for i in range(3)  # pages 0, 1, 2
        ]
        mock_doc = MagicMock()
        mock_doc.__iter__ = MagicMock(return_value=iter(pages))
        mock_pymupdf_open.return_value = mock_doc

        from docx import Document as RealDocument
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            out_path = f.name
        try:
            from pdf2zh.export_word import export_pdf_to_word
            export_pdf_to_word("fake.pdf", None, out_path, pages=[0, 2])

            doc = RealDocument(out_path)
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            self.assertIn("第0页内容", paras)
            self.assertNotIn("第1页内容", paras)  # page 1 not in pages list
            self.assertIn("第2页内容", paras)
        finally:
            os.unlink(out_path)

    @patch("pymupdf.open")
    def test_cjk_filter_strips_untranslated_english(self, mock_pymupdf_open):
        """Untranslated English blocks must be filtered when lang_out is Chinese."""
        mock_doc = MagicMock()
        mock_page = self._make_mock_page([
            {"x0": 50.0, "y0": 100.0, "x1": 500.0, "y1": 120.0, "content": "Untranslated English text"},
            {"x0": 50.0, "y0": 130.0, "x1": 500.0, "y1": 150.0, "content": self._ZH},
        ])
        mock_doc.__iter__ = MagicMock(return_value=iter([mock_page]))
        mock_pymupdf_open.return_value = mock_doc

        from docx import Document as RealDocument
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            out_path = f.name
        try:
            from pdf2zh.export_word import export_pdf_to_word
            export_pdf_to_word("fake.pdf", None, out_path, lang_out="zh")

            doc = RealDocument(out_path)
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            self.assertNotIn("Untranslated English text", paras)
            self.assertIn(self._ZH, paras)
        finally:
            os.unlink(out_path)


if __name__ == "__main__":
    unittest.main()
